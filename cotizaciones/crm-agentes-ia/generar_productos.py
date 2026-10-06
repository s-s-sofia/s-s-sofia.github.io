#!/usr/bin/env python3
"""Genera productos.js a partir de PRODUCTOS.xlsx para la demo del CRM.

Uso:  python3 generar_productos.py
Imágenes: guárdelas en img/productos/ con el código del artículo como nombre
(por ejemplo img/productos/500106.jpg) y vuelva a ejecutar este script.
"""
import json
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
IMG_DIR = BASE / "img" / "productos"
EXTS = (".webp", ".jpg", ".jpeg", ".png")

ws = openpyxl.load_workbook(BASE / "PRODUCTOS.xlsx", data_only=True).active
headers = [str(c.value).strip().upper() for c in ws[1]]
col = {name: headers.index(name) for name in ("LINEA", "ARTICULO", "ARTICULO_DESCRIPCION")}

productos = []
for row in ws.iter_rows(min_row=2, values_only=True):
    codigo = str(row[col["ARTICULO"]] or "").strip()
    if not codigo:
        continue
    img = next((f"img/productos/{codigo}{e}" for e in EXTS if (IMG_DIR / f"{codigo}{e}").exists()), None)
    productos.append({
        "linea": str(row[col["LINEA"]] or "").strip(),
        "codigo": codigo,
        "desc": " ".join(str(row[col["ARTICULO_DESCRIPCION"]] or "").split()),
        "img": img,
    })

out = BASE / "productos.js"
out.write_text(
    "// Generado por generar_productos.py a partir de PRODUCTOS.xlsx. No editar a mano.\n"
    f"window.SOFIA_PRODUCTOS = {json.dumps(productos, ensure_ascii=False, indent=1)};\n",
    encoding="utf-8",
)
con_img = sum(1 for p in productos if p["img"])
print(f"{len(productos)} productos ({con_img} con imagen) -> {out.name}")
